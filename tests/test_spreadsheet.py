"""Tests for the analysis report spreadsheet builders in common/spreadsheet.py."""

import io

import openpyxl
import pandas as pd

from common import spreadsheet as S
from fakes import FakeAnswer, trace

DF = pd.DataFrame({"Node": ["r1", "r2"], "Empty": [None, None], "Val": ["a", None]})


# ----------------------------------------------------------- _safe_sheet_name
def test_sheet_names_replace_forbidden_characters():
    assert S._safe_sheet_name("a/b:c", set()) == "a_b_c"
    assert S._safe_sheet_name("[x]?*", set()) == "_x___"
    assert S._safe_sheet_name("a\\b", set()) == "a_b"


def test_sheet_names_are_cut_to_31_characters():
    assert S._safe_sheet_name("x" * 40, set()) == "x" * 31


def test_sheet_names_stay_unique_and_within_the_limit():
    used = set()
    names = [S._safe_sheet_name(n, used) for n in ["same", "same", "x" * 40, "x" * 40]]
    assert names[:2] == ["same", "same_1"]
    assert len(set(names)) == 4 and all(len(n) <= 31 for n in names)
    assert names[3].endswith("_1")
    assert used == set(names)


# --------------------------------------------------------- _traces_to_dataframe
def test_traces_are_numbered_and_carry_their_disposition():
    df = S._traces_to_dataframe([trace("ACCEPTED"), trace("DENIED_IN", node="r2")])
    assert list(df.columns) == [
        "Trace",
        "Disposition",
        "Node",
        "Type",
        "Action",
        "Detail",
    ]
    assert df["Trace"].tolist() == [1, 2]
    assert df["Disposition"].tolist() == ["ACCEPTED", "DENIED_IN"]
    assert df["Node"].tolist() == ["r1", "r2"]


def test_no_traces_gives_an_empty_frame():
    assert S._traces_to_dataframe([]).shape == (0, 0)


# -------------------------------------------------------- question_to_dataframe
def test_empty_or_missing_answers_give_none():
    assert S.question_to_dataframe("nodeProperties", FakeAnswer(empty=True)) is None
    assert S.question_to_dataframe("nodeProperties", None) is None


def test_table_answers_lose_empty_columns():
    df = S.question_to_dataframe("nodeProperties", FakeAnswer(DF))
    assert list(df.columns) == ["Node", "Val"]


def test_traceroute_answers_become_trace_tables():
    answer = FakeAnswer(rows=[{"Traces": [trace(), trace()]}])
    df = S.question_to_dataframe("traceroute", answer)
    assert df["Trace"].tolist() == [1, 2]


def test_bidirectional_answers_are_tagged_with_a_direction():
    answer = FakeAnswer(
        rows=[
            {
                "Forward_Traces": [trace("ACCEPTED")],
                "Reverse_Traces": [trace("NO_ROUTE")],
            }
        ]
    )
    df = S.question_to_dataframe("bidirectionalTraceroute", answer)
    assert df["Direction"].tolist() == ["Forward", "Reverse"]
    assert df["Disposition"].tolist() == ["ACCEPTED", "NO_ROUTE"]


def test_test_filters_rows_are_flattened():
    answer = FakeAnswer(rows=[{"Node": "n1", "Trace": []}])
    df = S.question_to_dataframe("testFilters", answer)
    assert df["Node"].tolist() == ["n1"] and "Trace_Text" in df.columns


def test_a_broken_answer_gives_none_instead_of_raising():
    assert (
        S.question_to_dataframe("nodeProperties", FakeAnswer(RuntimeError("x"))) is None
    )


# ------------------------------------------------------------- build_report_xlsx
def item(name, answer, options=None, fun="nodeProperties"):
    return {"name": name, "fun": fun, "answer": answer, "options": options}


def sheets(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_xlsx_has_one_sheet_per_answer_with_its_rows():
    data = S.build_report_xlsx([item("Props", FakeAnswer(DF))])
    wb = sheets(data)
    assert wb.sheetnames == ["Props"]
    rows = [[c.value for c in r] for r in wb["Props"].iter_rows()]
    assert rows[0] == ["Node", "Val"] and rows[1] == ["r1", "a"]


def test_xlsx_skips_empty_and_failed_answers():
    data = S.build_report_xlsx(
        [
            item("Real", FakeAnswer(DF)),
            item("Empty", FakeAnswer(pd.DataFrame())),
            item("None", None),
            item("Broken", FakeAnswer(RuntimeError("x"))),
        ]
    )
    assert sheets(data).sheetnames == ["Real"]


def test_xlsx_is_none_when_there_is_nothing_to_write():
    assert S.build_report_xlsx([]) is None
    assert (
        S.build_report_xlsx([item("E", FakeAnswer(pd.DataFrame())), item("N", None)])
        is None
    )


def test_xlsx_ignores_repeats_but_keeps_different_options():
    data = S.build_report_xlsx(
        [
            item("Props", FakeAnswer(DF), None),
            item("Props", FakeAnswer(DF), None),  # exact repeat
            item("Props", FakeAnswer(DF), {"a": 1}),  # same name, new options
        ]
    )
    assert sheets(data).sheetnames == ["Props", "Props_1"]


def test_select_questions_are_cleaned_like_the_report_page():
    df = pd.DataFrame(
        {"File": ["a", "b"], "Status": ["PASSED", None], "Gone": [None, None]}
    )
    out = S.question_to_dataframe("fileParseStatus", FakeAnswer(df))
    assert list(out.columns) == ["File", "Status"]
    assert out["Status"].tolist() == ["PASSED", ""]  # blank, not missing


def test_routes_answers_are_cleaned_like_other_select_questions():
    df = pd.DataFrame({"Node": ["r1", "r2"], "Network": ["10.0.0.0/24", None]})
    out = S.question_to_dataframe("routes", FakeAnswer(df))
    assert out["Network"].tolist() == ["10.0.0.0/24", ""]
