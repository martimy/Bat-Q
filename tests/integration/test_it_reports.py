"""
The report builders fed with real pybatfish answers (the unit tests use
stand-ins for these).
"""

import io

import openpyxl
import pytest
from pybatfish.datamodel.flow import HeaderConstraints

from common import assertion_report as R
from common import htmlreport as H
from common import spreadsheet as S
from common.assertions import run_assertion_item
from it_helpers import host_ip

pytestmark = pytest.mark.integration


def answer(bf, fun, **options):
    return getattr(bf.q, fun)(**options).answer(snapshot="live")


@pytest.fixture(scope="module")
def web_flow():
    return dict(
        startLocation="as1border1",
        headers=HeaderConstraints(dstIps=host_ip("host1"), applications=["HTTP"]),
    )


# ------------------------------------------------------------ analysis report
def test_table_answer(bf):
    a = answer(bf, "nodeProperties")
    html = H.render_result_html("nodeProperties", a)
    assert "<table" in html and "as1border1" in html and "bq-error" not in html
    assert "as1border1" in set(S.question_to_dataframe("nodeProperties", a)["Node"])


def test_routes_answer(bf):
    a = answer(bf, "routes", nodes="as1border1")
    html = H.render_result_html("routes", a)
    assert "<th>Network</th>" in html and "bq-error" not in html


@pytest.mark.parametrize("fun", ["layer3Edges", "bgpEdges"])
def test_topology_answers_have_a_plot_of_the_devices(bf, fun):
    html = H.render_result_html(fun, answer(bf, fun))
    assert "plotly" in html.lower() and "as1border1" in html
    assert "bq-error" not in html


def test_traceroute_answer(bf, web_flow):
    a = answer(bf, "traceroute", **web_flow)
    dispositions = {str(t.disposition) for t in a.frame()["Traces"].iloc[0]}
    html = H.render_result_html("traceroute", a)
    assert "bq-error" not in html and "Trace" in html
    assert all(d in html for d in dispositions), dispositions
    df = S.question_to_dataframe("traceroute", a)
    assert dispositions <= set(df["Disposition"])
    assert {"Trace", "Node", "Type", "Action"} <= set(df.columns)


def test_test_filters_answer(bf):
    a = answer(
        bf,
        "testFilters",
        filters="OUTSIDE_TO_INSIDE",
        headers=HeaderConstraints(srcIps="1.0.0.1", dstIps=host_ip("host1")),
    )
    df = S.question_to_dataframe("testFilters", a)
    assert df is not None and {"Node", "Action", "Trace_Text"} <= set(df.columns)
    html = H.render_result_html("testFilters", a)
    assert "<th>Trace_Text</th>" in html and "bq-error" not in html


def test_the_whole_analysis_report_and_spreadsheet(bf, web_flow):
    calls = [
        ("Node Properties", "nodeProperties", {}),
        ("Layer 3 Edges", "layer3Edges", {}),
        ("Traceroute", "traceroute", web_flow),
    ]
    sections, data = [], []
    for name, fun, options in calls:
        a = answer(bf, fun, **options)
        sections.append(
            {"name": name, "options": None, "html": H.render_result_html(fun, a)}
        )
        data.append({"name": name, "fun": fun, "answer": a, "options": None})
    page = H.build_report_html(sections, "live", "now")
    assert page.count("<section") == 3 and "as1border1" in page
    wb = openpyxl.load_workbook(io.BytesIO(S.build_report_xlsx(data)))
    assert wb.sheetnames == ["Node Properties", "Layer 3 Edges", "Traceroute"]
    assert all(wb[n].max_row > 1 for n in wb.sheetnames)


# ----------------------------------------------------------- assertion report
def failing_item():
    return {
        "id": "a1",
        "type": "assert_zero_results",
        "name": "No nodes",
        "description": "",
        "scope": "active",
        "params": {},
        "question": {"fun": "nodeProperties", "options": {}},
    }


def test_assertion_report_shows_the_rows_of_a_failed_assertion(bf_state):
    item = failing_item()
    result = run_assertion_item(item, "live")
    assert result.status == "fail"
    rows = R.collect_report_rows([item], {"a1@live": result}, "live")
    html = R.build_assertion_report_html(rows, "live", "now")
    assert R.VERDICT_FAILED in html and "as1border1" in html
    wb = openpyxl.load_workbook(
        io.BytesIO(R.build_assertion_report_xlsx(rows, "live", "now"))
    )
    assert len(wb.sheetnames) == 2
    values = {c.value for row in wb[wb.sheetnames[1]].iter_rows() for c in row}
    assert "as1border1" in values


def test_assertion_report_with_a_real_flow_failure(bf_state):
    result = run_assertion_item(
        {
            "id": "a1",
            "type": "assert_flows_succeed",
            "params": {
                "startLocation": "no_such_device",
                "headers": {"dstIps": host_ip("host1")},
            },
        },
        "live",
    )
    item = {"id": "a1", "type": "assert_flows_succeed", "name": "Flow", "params": {}}
    rows = R.collect_report_rows([item], {"a1@live": result}, "live")
    html = R.build_assertion_report_html(rows, "live", "now")
    assert result.status == "error"
    assert R.VERDICT_INCONCLUSIVE in html and "could not be evaluated" in html
